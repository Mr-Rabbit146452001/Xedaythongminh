import os
import re
import urllib.parse
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_backend_docx():
    doc = docx.Document()
    
    # Page Setup: Standard Vietnamese Thesis Margins (A4: Top 2cm, Bottom 2cm, Left 3cm, Right 2cm)
    for s in doc.sections:
        s.top_margin = Inches(0.82)     # ~2.08 cm
        s.bottom_margin = Inches(0.82)  # ~2.08 cm
        s.left_margin = Inches(1.18)    # 3.0 cm
        s.right_margin = Inches(0.79)   # 2.0 cm
        
        # Configure Header & Footer
        footer = s.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("Đồ án Tốt nghiệp: Xe đẩy Mua sắm Thông minh AIoT — Chuyên đề Backend")
        r_ft.font.name = "Times New Roman"
        r_ft.font.size = Pt(9)
        r_ft.italic = True
        r_ft.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Set base font
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Times New Roman'
    font_normal.size = Pt(12)
    font_normal.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_cell_border(cell, **kwargs):
        tc = cell._element
        tcPr = tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
        for edge in ('top', 'left', 'bottom', 'right'):
            edge_data = kwargs.get(edge)
            if edge_data:
                tag = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", 4)}" w:space="0" w:color="{edge_data.get("color", "CBD5E1")}"/>'
                tcBorders.append(parse_xml(tag))
        tcPr.append(tcBorders)

    def add_formatted_runs(paragraph, text, base_font_size=12.0):
        clean_text = text.replace('$', '')
        tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', clean_text)
        for token in tokens:
            if not token:
                continue
            if token.startswith('**') and token.endswith('**') and len(token) >= 4:
                r = paragraph.add_run(token[2:-2])
                r.font.name = "Times New Roman"
                r.font.size = Pt(base_font_size)
                r.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
                r = paragraph.add_run(token[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(base_font_size - 1.5)
                r.bold = True
                r.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)
            elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
                r = paragraph.add_run(token[1:-1])
                r.font.name = "Times New Roman"
                r.font.size = Pt(base_font_size)
                r.italic = True
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            else:
                r = paragraph.add_run(token)
                r.font.name = "Times New Roman"
                r.font.size = Pt(base_font_size)
                r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(15)
        run.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11.5)
        run.italic = True
        run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13.5)
        run.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12.5)
        run.bold = True
        run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        return p

    def add_body(text, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.25
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_formatted_runs(p, text, base_font_size=12.0)
        return p

    def add_bullet(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.22
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_formatted_runs(p, text, base_font_size=12.0)
        return p

    def add_image_safe(img_path, caption_text, width_inches=5.8):
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run()
            run.add_picture(img_path, width=Inches(width_inches))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(10)
            run_cap = p_cap.add_run(caption_text)
            run_cap.font.name = "Times New Roman"
            run_cap.font.size = Pt(10)
            run_cap.italic = True
            run_cap.bold = True
            run_cap.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        else:
            p_err = doc.add_paragraph()
            p_err.add_run(f"[Hình ảnh chưa tìm thấy: {img_path}]").bold = True

    def add_code_block(code_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.3)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        border_spec = {"val": "single", "sz": 4, "color": "CBD5E1"}
        set_cell_border(cell, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
        
        p_cell = cell.paragraphs[0]
        p_cell.paragraph_format.space_before = Pt(0)
        p_cell.paragraph_format.space_after = Pt(0)
        run = p_cell.add_run(code_text.strip())
        run.font.name = "Consolas"
        run.font.size = Pt(8.8)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        doc.add_paragraph().paragraph_format.space_after = Pt(1)

    # =========================================================================
    # TRANG BÌA CHUYÊN ĐỀ ĐỒ ÁN TỐT NGHIỆP
    # =========================================================================
    add_title("BÁO CÁO KỸ THUẬT CHUYÊN ĐỀ ĐỒ ÁN TỐT NGHIỆP")
    add_title("MÔ HÌNH XE ĐẨY MUA SẮM THÔNG MINH TÍCH HỢP AIoT CHO SIÊU THỊ BÁN LẺ")
    add_subtitle("HỌC VIỆN / TRƯỜNG ĐẠI HỌC — NGÀNH KỸ THUẬT ĐIỆN TỬ - VIỄN THÔNG / CNTT")
    
    tb_summary = doc.add_table(rows=7, cols=2)
    tb_summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary_data = [
        ("Phân hệ chuyên trách:", "Thành viên 3 — Chuyên trách Backend (FastAPI / PostgreSQL / API Gateway / Mock Bank)"),
        ("Kiến trúc phân tầng (Mục 3.1):", "Kiến trúc 3 tầng: Biên IoT (Pi 4 8GB + 1 Cân HX711 tâm giỏ + YOLO11n + GM65) → Backend trung tâm → Giao diện"),
        ("Thiết kế CSDL & API (Mục 3.5.1):", "CSDL PostgreSQL 7 bảng chuẩn hóa 3NF, bảo đảm ACID, khóa bất biến + 6 REST API Endpoints FastAPI"),
        ("Kiểm thử kịch bản (Mục 4.1):", "Bộ 230 lượt kiểm thử tích hợp 5 ca kiểm thử chức năng có chủ đích (Happy Path, Lệch cân, AI mờ, GM65, Rớt mạng)"),
        ("Tỷ lệ thành công (Mục 4.4):", "Bộ 150 cặp thử (300 lượt vào/ra) trên 5 sản phẩm cốt lõi: Hệ thống đúng 100.0% (89.67% Camera AI + 10.33% GM65 bổ sung)"),
        ("Hiệu năng Backend (Mục 4.5.1):", "Đo lường 500 yêu cầu đối soát: Độ trễ TB 28.4ms, P50 24.8ms, P99 58.1ms (vượt chuẩn SLA 50ms & 100ms)"),
        ("Trạng thái nghiệm thu:", "HOÀN THÀNH 100% CÁC TIÊU CHÍ KỸ THUẬT VÀ PHẢN BIỆN (14/14 MỤC)")
    ]
    
    for r_idx, (k, v) in enumerate(summary_data):
        row = tb_summary.rows[r_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.1)
        c0.text = k
        c1.text = v
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")
        set_cell_margins(c0, top=60, bottom=60, left=90, right=90)
        set_cell_margins(c1, top=60, bottom=60, left=90, right=90)
        b_spec = {"val": "single", "sz": 4, "color": "CBD5E1"}
        set_cell_border(c0, top=b_spec, bottom=b_spec, left=b_spec, right=b_spec)
        set_cell_border(c1, top=b_spec, bottom=b_spec, left=b_spec, right=b_spec)
        c0.paragraphs[0].runs[0].font.name = "Times New Roman"
        c0.paragraphs[0].runs[0].font.size = Pt(10)
        c0.paragraphs[0].runs[0].bold = True
        c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        c1.paragraphs[0].runs[0].font.name = "Times New Roman"
        c1.paragraphs[0].runs[0].font.size = Pt(10)
        if r_idx == 6:
            c1.paragraphs[0].runs[0].bold = True
            c1.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x04, 0x78, 0x57)

    doc.add_page_break()

    # =========================================================================
    # ĐỌC NỘI DUNG TỪ MARKDOWN VÀ DỰNG TRANG
    # =========================================================================
    md_path = r"c:\Users\LE THI HAU\.gemini\antigravity\scratch\StrollerApp\BaoCao_Backend_HeThong_Chuong3_4.md"
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    lines = md_text.split('\n')
    i = 0
    in_code_block = False
    code_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Bỏ qua phần tiêu đề ban đầu trong MD vì đã có bìa
        if i < 4 and (stripped.startswith('#') or stripped.startswith('**')):
            i += 1
            continue

        # Xử lý Code Block (``` ... ```)
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_buffer = []
            else:
                in_code_block = False
                add_code_block("\n".join(code_buffer))
                code_buffer = []
            i += 1
            continue

        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Tiêu đề cấp 1 (# ...)
        if stripped.startswith('# '):
            add_heading_1(stripped[2:].strip())
            i += 1
            continue

        # Tiêu đề cấp 2 (## ...)
        if stripped.startswith('## '):
            add_heading_1(stripped[3:].strip())
            i += 1
            continue

        # Tiêu đề cấp 3 (### ...)
        if stripped.startswith('### '):
            add_heading_2(stripped[4:].strip())
            i += 1
            continue

        # Tiêu đề cấp 4 (#### ...)
        if stripped.startswith('#### '):
            add_heading_3(stripped[5:].strip())
            i += 1
            continue

        # Tiêu đề cấp 5/6 (##### hoặc ######)
        if stripped.startswith('##### ') or stripped.startswith('###### '):
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(6)
            p_sub.paragraph_format.space_after = Pt(2)
            p_sub.paragraph_format.keep_with_next = True
            r_sub = p_sub.add_run(stripped.lstrip('#').strip())
            r_sub.font.name = "Times New Roman"
            r_sub.font.size = Pt(11.5)
            r_sub.bold = True
            r_sub.italic = True
            r_sub.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        # Danh sách Bullet (- ...)
        if stripped.startswith('- ') or stripped.startswith('* '):
            bullet_content = stripped[2:].strip()
            add_bullet(bullet_content)
            i += 1
            continue

        # Bảng dữ liệu Markdown (| ... |)
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 2:
                headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
                data_rows = []
                for t_line in table_lines[2:]:
                    cols = [c.strip() for c in t_line.split('|')[1:-1]]
                    data_rows.append(cols)

                doc_table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # Format Header Row
                hdr_cells = doc_table.rows[0].cells
                for col_idx, h_text in enumerate(headers):
                    cell = hdr_cells[col_idx]
                    cell.text = h_text.replace('**', '').replace('$', '')
                    set_cell_background(cell, "1E3A8A")
                    set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
                    border_spec = {"val": "single", "sz": 4, "color": "CBD5E1"}
                    set_cell_border(cell, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
                    p_hdr = cell.paragraphs[0]
                    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p_hdr.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(9.0)
                        r.bold = True
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

                # Format Data Rows
                for r_idx, r_data in enumerate(data_rows):
                    row_cells = doc_table.rows[r_idx + 1].cells
                    base_bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    for col_idx in range(len(headers)):
                        val = r_data[col_idx] if col_idx < len(r_data) else ""
                        cell = row_cells[col_idx]
                        cell.text = ""
                        p_d = cell.paragraphs[0]
                        p_d.paragraph_format.space_before = Pt(0)
                        p_d.paragraph_format.space_after = Pt(0)
                        add_formatted_runs(p_d, val, base_font_size=8.8)
                        
                        # Tô màu vàng nhạt cho cột GM65 bổ sung đúng giống như bảng chuẩn của nhóm nghiên cứu
                        if "GM65 bổ sung" in headers[col_idx]:
                            cell_bg = "FEF9C3" if r_idx % 2 == 0 else "FEF08A"
                        else:
                            cell_bg = base_bg

                        set_cell_background(cell, cell_bg)
                        set_cell_margins(cell, top=55, bottom=55, left=80, right=80)
                        border_spec = {"val": "single", "sz": 4, "color": "E2E8F0"}
                        set_cell_border(cell, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
                        if col_idx in [0, 1, 2, 3, 4, 5, 6, 7]:
                            # Căn giữa các cột số liệu hoặc tiêu đề ngắn
                            if col_idx == 0 and len(val) > 15:
                                p_d.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            else:
                                p_d.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        else:
                            p_d.alignment = WD_ALIGN_PARAGRAPH.LEFT

                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # Xử lý hình ảnh Markdown (![...](...))
        img_match = re.search(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            caption = img_match.group(1).strip()
            raw_path = img_match.group(2).strip()
            clean_path = urllib.parse.unquote(raw_path.replace('file:///', '').replace('file://', '')).replace('/', '\\')
            
            # Đọc caption ở dòng tiếp theo nếu có cú pháp "*Hình x.x: ...*"
            if i + 1 < len(lines) and lines[i+1].strip().startswith('*Hình'):
                caption = lines[i+1].strip().replace('*', '')
                i += 1

            add_image_safe(clean_path, caption)
            i += 1
            continue

        # Xử lý Blockquote (> ...)
        if stripped.startswith('> '):
            quote_text = stripped[2:].strip().replace('**', '')
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(6)
            p_q.paragraph_format.left_indent = Inches(0.4)
            p_q.paragraph_format.right_indent = Inches(0.4)
            run_q = p_q.add_run(quote_text)
            run_q.font.name = "Times New Roman"
            run_q.font.size = Pt(11)
            run_q.italic = True
            run_q.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        # Dòng phân cách (---)
        if stripped == '---':
            i += 1
            continue

        # Đoạn văn thông thường
        add_body(stripped)
        i += 1

    # =========================================================================
    # LƯU FILE WORD VÀO TẤT CẢ CÁC ĐÍCH ĐẾN
    # =========================================================================
    destinations = [
        r"c:\Users\LE THI HAU\.gemini\antigravity\scratch\StrollerApp\BaoCao_Backend_HeThong_Chuong3_4_CapNhat_GM65.docx",
        r"c:\Users\LE THI HAU\.gemini\antigravity\scratch\StrollerApp\BaoCao_Backend_HeThong_Chuong3_4_Final.docx",
        r"c:\Users\LE THI HAU\.gemini\antigravity\scratch\StrollerApp\BaoCao_Backend_HeThong_Chuong3_4.docx",
        r"C:\Users\LE THI HAU\Pictures\Screenshots\BaoCao_Backend_HeThong_Chuong3_4.docx",
        r"C:\Users\LE THI HAU\Downloads\BaoCao_Backend_HeThong_Chuong3_4.docx",
        r"c:\Users\LE THI HAU\Downloads\cơ sơ dữ liệu-0\BaoCao_Backend_HeThong_Chuong3_4.docx"
    ]

    for dest in destinations:
        try:
            os.makedirs(os.path.dirname(dest), exist_ok=True)
            doc.save(dest)
            print(f"✅ Đã lưu file Word thành công tại: {dest}")
        except Exception as e:
            print(f"⚠️ Không thể lưu tại {dest}: {e}")

if __name__ == "__main__":
    create_backend_docx()
