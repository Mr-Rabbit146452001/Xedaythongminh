import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

doc = docx.Document()

# Page Setup: Normal Margins (1 inch)
sections = doc.sections
for s in sections:
    s.top_margin = Inches(1)
    s.bottom_margin = Inches(1)
    s.left_margin = Inches(1)
    s.right_margin = Inches(1)

# Helper function to set table cell background color
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_p.add_run("BÁO CÁO CHUYÊN ĐỀ ĐỒ ÁN TỐT NGHIỆP\n")
title_run.font.name = "Times New Roman"
title_run.font.size = Pt(14)
title_run.bold = True
title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

sub_title_run = title_p.add_run("THIẾT KẾ VÀ THI CÔNG MÔ HÌNH XE ĐẨY MUA SẮM THÔNG MINH TÍCH HỢP AIoT CHO SIÊU THỊ BÁN LẺ\n")
sub_title_run.font.name = "Times New Roman"
sub_title_run.font.size = Pt(16)
sub_title_run.bold = True
sub_title_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

part_title_run = title_p.add_run("PHÂN HỆ GIAO DIỆN VÀ TRẢI NGHIỆM NGƯỜI DÙNG (FRONTEND & UI/UX)\n(App Tablet Xe Đẩy, Web Admin Quản Trị Bán Lẻ & Web Khách Hàng PWA)")
part_title_run.font.name = "Times New Roman"
part_title_run.font.size = Pt(13)
part_title_run.italic = True
part_title_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

doc.add_paragraph()

# Author Info Table
info_table = doc.add_table(rows=4, cols=2)
info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
info_data = [
    ("Tên đề tài:", "Thiết kế và thi công mô hình xe đẩy mua sắm thông minh tích hợp AIoT cho siêu thị bán lẻ"),
    ("Thành viên thực hiện:", "Thành viên 4 – Chuyên trách App / Web Frontend & Thiết kế UI/UX"),
    ("Giảng viên hướng dẫn:", "ThS. Đỗ Văn Việt Em"),
    ("Thời gian thực hiện:", "Học kỳ 1 – Năm học 2026")
]
for i, (k, v) in enumerate(info_data):
    row = info_table.rows[i]
    cell_k, cell_v = row.cells[0], row.cells[1]
    cell_k.text = k
    cell_v.text = v
    cell_k.paragraphs[0].runs[0].bold = True
    cell_k.paragraphs[0].runs[0].font.name = "Times New Roman"
    cell_v.paragraphs[0].runs[0].font.name = "Times New Roman"
    set_cell_background(cell_k, "F1F5F9")
    set_cell_background(cell_v, "FFFFFF")

doc.add_paragraph()

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p

def add_heading_3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.bold = True
    run.italic = True
    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_body(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return p

def add_bullet(text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = "Times New Roman"
        r_bold.font.size = Pt(12)
        r_bold.bold = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return p

def add_image_caption(caption_text, file_ref):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(f"{caption_text}\n(Nguồn minh chứng: {file_ref})")
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.italic = True
    run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

# Save empty doc to begin
doc.save(r"c:\Users\LE THI HAU\.gemini\antigravity\scratch\StrollerApp\BaoCao_ToanDien_UI_UX_Frontend.docx")
print("Initialized docx")
